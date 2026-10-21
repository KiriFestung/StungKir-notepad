# StungKir - писательский блокнот. Copyright (C) 2026 KiriFesTung
# SPDX-License-Identifier: GPL-3.0-or-later
"""Панель «рядом» и шаблоны сборки."""
import unittest
import zipfile

from stungkir import config as C
from stungkir import export as E
from stungkir.tests.helpers import Sandbox, WindowCase, app, wait

CHAPTERS = [("Глава 1", ["Первый абзац.", "Второй абзац с *курсивом*.", "* * *", "После разрыва."]),
            ("Глава 2", ["Текст второй главы."])]


def template(name):
    look = {k: v for k, v in E.TEMPLATES[name].items() if k in E.LAYOUT}
    look["contacts"] = "mira@example.com"
    return {"headings": True, "new_page": True, "layout": look}


class TemplatesTest(Sandbox):
    def test_volume_line(self):
        self.assertEqual(E.book_volume(CHAPTERS), (70, 11))     # разрыв сцены не считается
        self.assertIn("70 знаков с пробелами", E.volume_line(CHAPTERS))
        self.assertIn("11 слов", E.volume_line(CHAPTERS))

    def test_scene_breaks(self):
        for line in ("***", "* * *", "#", "~ ~ ~", "•••"):
            self.assertTrue(E.is_scene_break(line), line)
        self.assertFalse(E.is_scene_break("**жирный**"))

    @unittest.skipUnless(E.docx_available(), "нет python-docx")
    def test_manuscript_docx(self):
        import docx
        target = self.tmp / "ms.docx"
        E.export_docx(target, "Книга", "Мира Лосс", CHAPTERS, template("Рукопись для издательства"))
        document = docx.Document(str(target))
        texts = [p.text for p in document.paragraphs]
        self.assertTrue(any("70 знаков с пробелами" in t for t in texts))
        self.assertIn("Мира Лосс\nmira@example.com", texts)
        self.assertEqual(document.styles["Normal"].paragraph_format.line_spacing, 1.5)
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.mm, 210, delta=0.5)      # A4, а не американский Letter
        self.assertAlmostEqual(section.left_margin.mm, 25, delta=0.5)
        header = section.header.paragraphs[0].text
        self.assertIn("Мира Лосс · Книга", header)
        scene = next(p for p in document.paragraphs if p.text == "* * *")
        self.assertEqual(scene.alignment, 1)          # по центру

    def test_print_pdf_page_size(self):
        app()
        target = self.tmp / "print.pdf"
        E.export_pdf(target, "Книга", "Автор", CHAPTERS, template("Печать, книжный формат"))
        data = target.read_bytes()
        self.assertTrue(data.startswith(b"%PDF"))
        import re
        box = re.search(rb"/MediaBox \[\s*0 0 ([\d.]+) ([\d.]+)", data)
        width_mm, height_mm = (float(box.group(i)) / 72 * 25.4 for i in (1, 2))
        self.assertAlmostEqual(width_mm, 148, delta=1)
        self.assertAlmostEqual(height_mm, 210, delta=1)

    def test_pdf_text_is_black_in_any_theme(self):
        """Текст книги - чёрный, даже если в программе зелёная тема (раньше цвет темы попадал в PDF)."""
        app()
        try:
            from PySide6.QtCore import QSize
            from PySide6.QtPdf import QPdfDocument
        except ImportError:
            self.skipTest("нет модуля QtPdf")
        from PySide6.QtWidgets import QApplication
        from stungkir.style import apply_theme
        C.set_palette(C.THEMES["Мох"])
        apply_theme(QApplication.instance())
        target = self.tmp / "book.pdf"
        E.export_pdf(target, "Книга", "Автор", CHAPTERS, template("Обычная книга"))
        pdf = QPdfDocument(None)
        pdf.load(str(target))
        image = pdf.render(1, QSize(600, 850))
        darkest = min(sum(image.pixelColor(x, y).getRgb()[:3])     # без нижнего поля с номером страницы
                      for x in range(0, image.width(), 2) for y in range(0, int(image.height() * 0.8), 2))
        self.assertLess(darkest, 60)          # есть по-настоящему чёрные буквы, а не бледно-зелёные
        pdf.close()

    def test_fb2_structure(self):
        import xml.etree.ElementTree as ET
        target = self.tmp / "book.fb2"
        chapters = [("Глава 1", ["# Пролог", "Абзац с *курсивом* и **жирным** & <скобками>.", "* * *", "После."]),
                    ("Пустая", [])]
        options = {**template("Самиздат FB2"), "genre": "sf_fantasy"}
        E.export_fb2(target, "Туман", "Ольга Петрова", chapters, options)
        ns = {"f": "http://www.gribuser.ru/xml/fictionbook/2.0"}
        root = ET.parse(target).getroot()                    # корректный XML
        self.assertEqual(root.find("f:description/f:title-info/f:genre", ns).text, "sf_fantasy")
        self.assertEqual(root.find("f:description/f:title-info/f:lang", ns).text, "ru")
        self.assertEqual(root.find("f:description/f:title-info/f:author/f:last-name", ns).text, "Петрова")
        sections = root.findall("f:body/f:section", ns)
        self.assertEqual(len(sections), 2)
        first = sections[0]
        self.assertEqual(first.find("f:title/f:p", ns).text, "Глава 1")
        self.assertEqual([s.text for s in first.findall("f:subtitle", ns)], ["Пролог", "* * *"])
        paragraph = first.findall("f:p", ns)[0]
        self.assertEqual(paragraph.find("f:emphasis", ns).text, "курсивом")
        self.assertIn("<скобками>", "".join(paragraph.itertext()))
        self.assertTrue(len(sections[1]) >= 2)              # пустая глава - не пустая секция
        E.export_fb2(self.tmp / "nick.fb2", "Т", "KiriFesTung", chapters, options)
        root = ET.parse(self.tmp / "nick.fb2").getroot()
        self.assertEqual(root.find("f:description/f:title-info/f:author/f:nickname", ns).text, "KiriFesTung")

    def test_epub_scene_break_and_indent(self):
        target = self.tmp / "b.epub"
        E.export_epub(target, "Книга", "", CHAPTERS, template("Самиздат EPUB"))
        with zipfile.ZipFile(target) as book:
            chapter = book.read("OEBPS/chapter-001.xhtml").decode()
            css = book.read("OEBPS/style.css").decode()
        self.assertIn('<p class="scene">* * *</p>', chapter)
        self.assertIn('<p class="first">После разрыва.</p>', chapter)
        self.assertIn("text-indent: 1.40em", css)


class CompileDialogTest(WindowCase):
    def test_template_sets_format_and_layout(self):
        from stungkir.windows import CompileDialog
        self.chapter("Книга", "Глава", "Текст.")
        w = self.window()
        dialog = CompileDialog(w, w.categories(), None, "Автор")
        dialog.template.setCurrentText("Рукопись для издательства")
        dialog.contacts.setText("mira@example.com")
        options = dialog.options()
        self.assertEqual(options["format"], "docx")
        self.assertEqual(options["layout"]["title_page"], "manuscript")
        self.assertEqual(options["layout"]["contacts"], "mira@example.com")
        again = CompileDialog(w, w.categories(), None, "Автор")      # выбор запоминается
        self.assertEqual(again.template.currentText(), "Рукопись для издательства")
        self.assertEqual(again.contacts.text(), "mira@example.com")


class SidePaneTest(WindowCase):
    def setUp(self):
        super().setUp()
        self.first = self.chapter("Книга", "Глава 1", "Первая глава.")
        self.second = self.chapter("Книга", "Глава 2", "Вторая глава.")
        world = C.ROOT / "Мир"
        world.mkdir()
        (world / "Тервена.txt").write_text("Прибрежное государство.", encoding="utf-8", newline="\n")
        self.w = self.window()
        self.w.open_path(self.first)

    def test_open_beside_and_swap(self):
        self.w.open_beside(self.second)
        wait()
        self.assertTrue(self.w.side.isVisible())
        self.assertEqual(self.w.side.view.toPlainText(), "Вторая глава.")
        self.assertTrue(self.w.side.view.isReadOnly())
        self.w.open_from_side(self.second)
        self.assertEqual(self.w.path, self.second)
        self.assertEqual(self.w.side.path, self.first)

    def test_clock_does_not_cover_side_pane(self):
        w = self.window()
        w.resize(1200, 700)
        w.show()
        wait(0.1)
        self.assertTrue(w.corner_clock.isVisible())              # часы в правом углу
        w.open_beside(self.second)                                 # панель «рядом» - справа
        wait(0.1)
        self.assertFalse(w.corner_clock.isVisible())
        w.close_side()
        self.assertTrue(w.corner_clock.isVisible())

    def test_side_follows_save_and_rename(self):
        self.w.open_beside(self.first)
        self.w.editor.setPlainText("Новый текст.")
        self.w.save()
        self.assertEqual(self.w.side.view.toPlainText(), "Новый текст.")
        self.patch("ask_text", lambda *a, **k: ("Первая", True))
        self.w.rename(self.first)
        self.assertEqual(self.w.side.path.name, "Первая.txt")

    def test_ctrl_click_on_term_opens_beside(self):
        self.w.toggle_lore_category(C.ROOT / "Мир")
        self.w.editor.setPlainText("Она вспомнила Тервену.")
        term = self.w.lore.terms[0]
        self.w.open_lore_term(term)
        self.assertEqual(self.w.path, self.first)                  # глава осталась на месте
        self.assertEqual(self.w.side.path.name, "Тервена.txt")

    def test_toggle_and_table(self):
        table = C.ROOT / "Мир" / "Ресурсы.csv"
        table.write_text("Название,Описание\nЛунит,Светится\n", encoding="utf-8", newline="\n")
        self.w.open_beside(table)
        self.assertIn("Лунит", self.w.side.view.toPlainText())
        self.w.toggle_side()
        self.assertFalse(self.w.side.isVisible())
        self.w.toggle_side()
        self.assertEqual(self.w.side.path, table)
