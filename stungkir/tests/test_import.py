# StungKir - писательский блокнот. Copyright (C) 2026 KiriFesTung
# SPDX-License-Identifier: GPL-3.0-or-later
"""Импорт: Word (с разрезкой на главы), RTF в трёх видах, проекты Scrivener."""
import json
import unittest

from stungkir import config as C
from stungkir import importers as I
from stungkir import storage as S
from stungkir.export import docx_available
from stungkir.tests.helpers import Sandbox, WindowCase


def make_docx(path, headings=True):
    import docx
    d = docx.Document()
    if headings:
        d.add_paragraph("Вступление до первой главы.")
        d.add_heading("Глава 1. Туман", level=1)
    else:
        d.add_paragraph("ГЛАВА ПЕРВАЯ")
    p = d.add_paragraph("Обычный текст, ")
    p.add_run("курсив").italic = True
    p.add_run(" и ")
    p.add_run("жирный").bold = True
    p.add_run(".")
    d.add_paragraph("")
    d.add_paragraph("— Реплика!")
    if headings:
        d.add_heading("Подзаголовок", level=2)
        d.add_heading("Глава 2: «Ночь»/конец", level=1)
    else:
        d.add_paragraph("Глава 2")
    d.add_paragraph("Вторая глава.")
    d.save(str(path))


class HelpersTest(unittest.TestCase):
    def test_runs_to_markup(self):
        runs = [("Это ", False, False), ("кур", False, True), ("сив", False, True), (" и ", False, False),
                ("жир ", True, False), ("оба", True, True)]
        self.assertEqual(I.runs_to_markup(runs), "Это *курсив* и **жир** ***оба***")

    def test_safe_name(self):
        self.assertEqual(I.safe_name('Глава 2: «Ночь»/конец?'), "Глава 2- «Ночь»-конец")
        self.assertEqual(len(I.safe_name("очень " * 40)) <= I.NAME_LIMIT, True)
        self.assertEqual(I.safe_name("   ", "Запас"), "Запас")

    def test_split_by_chapter_lines_without_styles(self):
        paragraphs = [("Пролог", 0), ("Текст пролога.", 0), ("Глава 1", 0), ("Текст.", 0),
                      ("В главе было много слов и это обычная длинная строка текста про главу.", 0)]
        chapters = I.split_chapters(paragraphs, "Книга")
        self.assertEqual([name for name, _ in chapters], ["Пролог", "Глава 1"])
        self.assertEqual(len(chapters[1][1]), 2)


@unittest.skipUnless(docx_available(), "нет python-docx")
class DocxTest(Sandbox):
    def test_paragraphs_keep_markup_and_levels(self):
        make_docx(self.tmp / "book.docx")
        paragraphs = I.docx_paragraphs(self.tmp / "book.docx")
        self.assertIn(("Глава 1. Туман", 1), paragraphs)
        self.assertIn(("Обычный текст, *курсив* и **жирный**.", 0), paragraphs)
        chapters = I.split_chapters(paragraphs, "book")
        self.assertEqual([name for name, _ in chapters], ["book", "Глава 1. Туман", "Глава 2: «Ночь»/конец"])
        self.assertIn("Подзаголовок", chapters[1][1])          # заголовок пониже - строка текста
        self.assertNotIn("", chapters[1][1])                    # пустые абзацы выброшены

    def test_without_heading_styles_uses_chapter_words(self):
        make_docx(self.tmp / "plain.docx", headings=False)
        chapters = I.split_chapters(I.docx_paragraphs(self.tmp / "plain.docx"), "plain")
        self.assertEqual([name for name, _ in chapters], ["ГЛАВА ПЕРВАЯ", "Глава 2"])


@unittest.skipUnless(docx_available(), "нет python-docx")
class WordWindowTest(WindowCase):
    def setUp(self):
        super().setUp()
        self.book = C.ROOT / "Книга"
        self.book.mkdir()
        make_docx(self.tmp / "Роман.docx")
        self.w = self.window()
        self.w.tree.setCurrentItem(self.w.tree.topLevelItem(0))
        self.patch("QFileDialog", type("Files", (), {
            "getOpenFileNames": staticmethod(lambda *a, **k: ([str(self.tmp / "Роман.docx")], ""))}))

    def test_split_into_chapters(self):
        self.patch("ask_choice", lambda *a, **k: "split")
        self.w.import_files()
        self.assertEqual([f.name for f in S.files_in(self.book)],
                         ["Роман.txt", "Глава 1. Туман.txt", "Глава 2- «Ночь»-конец.txt"])
        self.assertEqual(S.read_text(self.book / "Глава 1. Туман.txt"),
                         "Обычный текст, *курсив* и **жирный**.\n— Реплика!\nПодзаголовок")

    def test_one_file(self):
        self.patch("ask_choice", lambda *a, **k: "one")
        self.w.import_files()
        text = S.read_text(self.book / "Роман.txt")
        self.assertTrue(text.startswith("Вступление до первой главы.\nГлава 1. Туман\n"))

    def test_broken_docx_is_reported(self):
        (self.tmp / "Роман.docx").write_bytes(b"not a zip")
        self.w.import_files()
        self.assertIn("не читается как файл Word", self.messages[-1][1])
        self.assertEqual(S.files_in(self.book), [])


class RtfTest(unittest.TestCase):
    def test_windows_codepage(self):
        rtf = (b"{\\rtf1\\ansi\\ansicpg1251{\\fonttbl{\\f0 Times;}}\\pard \\'cf\\'f0\\'e8\\'e2\\'e5\\'f2, "
               b"{\\i \\'ec\\'e8\\'f0}!\\par {\\b\\'c6}\\'e8\\'f0\\par}")
        self.assertEqual(I.rtf_to_lines(rtf), ["Привет, *мир*!", "**Ж**ир"])

    def test_mac_unicode_and_backslash_newline(self):
        rtf = (b"{\\rtf1\\ansi\\ansicpg1252\\cocoartf2639{\\fonttbl\\f0\\fnil Cochin;}{\\*\\expandedcolortbl;;}\n"
               b"\\pard \\uc0\\u1058 \\u1091 \\u1084 \\u1072 \\u1085  \\i \\uc0\\u1085 \\u1072 \\u1076 \\i0  "
               b"\\uc0\\u1088 \\u1077 \\u1082 \\u1086 \\u1081 .\\\n\\\n\\b \\uc0\\u1042 \\u1090 \\u1086 "
               b"\\u1088 \\u1086 \\u1081 \\b0  \\uc0\\u1072 \\u1073 \\u1079 \\u1072 \\u1094 .}")
        self.assertEqual(I.rtf_to_lines(rtf), ["Туман *над* рекой.", "", "**Второй** абзац."])

    def test_word_style_unicode_with_fallback(self):
        rtf = (b"{\\rtf1\\ansi\\uc1{\\stylesheet{\\s1 \\b Heading;}}{\\info{\\title X}}"
               b"\\pard \\u1045?\\u1083?\\u1082?\\u1072? \\emdash  {\\field{\\*\\fldinst HYPERLINK x}"
               b"{\\fldrslt \\u1089?\\u1089\\'3f\\u1099?\\u1083?\\u1082?\\u1072?}}\\par}")
        self.assertEqual(I.rtf_to_lines(rtf), ["Елка — ссылка"])


SCRIVX = """<?xml version="1.0" encoding="UTF-8"?>
<ScrivenerProject Version="2.0">
<Binder>
 <BinderItem UUID="D" Type="DraftFolder"><Title>Рукопись</Title><Children>
  <BinderItem UUID="CH1" Type="Folder"><Title>Глава 1</Title><Children>
   <BinderItem UUID="S1" Type="Text"><Title>Утро</Title></BinderItem>
   <BinderItem UUID="S2" Type="Text"><Title>Река</Title></BinderItem>
  </Children></BinderItem>
  <BinderItem UUID="S3" Type="Text"><Title>Эпилог: конец?</Title></BinderItem>
 </Children></BinderItem>
 <BinderItem UUID="R" Type="ResearchFolder"><Title>Материалы</Title><Children>
  <BinderItem UUID="P1" Type="PDF"><Title>Карта</Title></BinderItem>
  <BinderItem UUID="N1" Type="Text"><Title>Заметки о мире</Title></BinderItem>
 </Children></BinderItem>
 <BinderItem UUID="CHAR" Type="Folder"><Title>Персонажи</Title><Children>
  <BinderItem UUID="O" Type="Text"><Title>Ольга</Title></BinderItem>
 </Children></BinderItem>
 <BinderItem UUID="T" Type="TrashFolder"><Title>Корзина</Title><Children>
  <BinderItem UUID="X" Type="Text"><Title>Удалённое</Title></BinderItem>
 </Children></BinderItem>
</Binder>
</ScrivenerProject>
"""


def rtf(text):
    body = "".join(f"\\u{ord(ch)}?" if ord(ch) > 127 else ch for ch in text).replace("\n", "\\par ")
    return ("{\\rtf1\\ansi\\uc1 " + body + "}").encode("ascii")


def make_scriv(folder):
    folder.mkdir()
    (folder / "Книга.scrivx").write_text(SCRIVX, encoding="utf-8", newline="\n")
    texts = {"S1": "Утро было серым.\nТуман.", "S2": "Река текла.", "S3": "Конец.", "N1": "Мир круглый.",
             "O": "Рыжая и смелая.", "X": "Не надо."}
    for ident, text in texts.items():
        data = folder / "Files" / "Data" / ident
        data.mkdir(parents=True)
        (data / "content.rtf").write_bytes(rtf(text))
    (folder / "Files" / "Data" / "S2" / "synopsis.txt").write_text("Героиня у реки.", encoding="utf-8", newline="\n")
    (folder / "Files" / "Data" / "CH1").mkdir()
    (folder / "Files" / "Data" / "CH1" / "notes.rtf").write_bytes(rtf("Проверить сезон."))


class ScrivenerTest(Sandbox):
    def test_plan(self):
        make_scriv(self.tmp / "Книга.scriv")
        name, items = I.read_scrivener(self.tmp / "Книга.scriv")
        plan, skipped = I.scrivener_plan(name, items)
        self.assertEqual(skipped, 1)       # PDF
        plan = dict(plan)
        self.assertEqual(sorted(plan), ["Книга", "Книга · материалы", "Персонажи"])
        chapters = {title: (lines, notes) for title, lines, notes in plan["Книга"]}
        self.assertEqual(list(chapters), ["Глава 1", "Эпилог- конец"])
        lines, notes = chapters["Глава 1"]
        self.assertEqual(lines, ["Утро было серым.", "Туман.", "* * *", "Река текла."])
        self.assertEqual(notes, [(0, "Глава 1: Проверить сезон."), (3, "Река: Героиня у реки.")])

    def test_not_a_project(self):
        (self.tmp / "пусто").mkdir()
        with self.assertRaises(I.ScrivenerError):
            I.read_scrivener(self.tmp / "пусто")


class ScrivenerWindowTest(WindowCase):
    def test_import_creates_categories_and_notes(self):
        make_scriv(self.tmp / "Книга.scriv")
        w = self.window()
        self.patch("QFileDialog", type("Files", (), {
            "getExistingDirectory": staticmethod(lambda *a, **k: str(self.tmp / "Книга.scriv"))}))
        self.patch("ask_choice", lambda *a, **k: "go")
        w.import_scrivener()
        chapter = C.ROOT / "Книга" / "Глава 1.txt"
        self.assertEqual(S.read_text(chapter), "Утро было серым.\nТуман.\n* * *\nРека текла.")
        notes = json.loads(S.read_text(S.notes_path(chapter)))
        self.assertEqual([(n["anchor"], n["text"]) for n in notes],
                         [("Утро было серым.", "Глава 1: Проверить сезон."), ("Река текла.", "Река: Героиня у реки.")])
        self.assertTrue((C.ROOT / "Персонажи" / "Ольга.txt").exists())
        self.assertIn("Пропущено 1 файл", self.messages[-1][1])
        w.open_path(chapter)
        self.assertEqual(len(w.notes_margin.notes), 2)


ODT_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
<office:automatic-styles>
 <style:style style:name="T1" style:family="text"><style:text-properties fo:font-style="italic"/></style:style>
 <style:style style:name="T2" style:family="text"><style:text-properties fo:font-weight="bold"/></style:style>
</office:automatic-styles>
<office:body><office:text>
 <text:h text:outline-level="1">Глава 1. Туман</text:h>
 <text:p>Обычный <text:span text:style-name="T1">курсив</text:span> и <text:span text:style-name="T2">жирный</text:span>.</text:p>
 <text:h text:outline-level="1">Глава 2</text:h>
 <text:p>Строка<text:line-break/>вторая</text:p>
</office:text></office:body></office:document-content>"""

ODS_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
<office:body><office:spreadsheet><table:table table:name="Бестиарий">
 <table:table-row><table:table-cell><text:p>Существо</text:p></table:table-cell>
  <table:table-cell><text:p>Где водится</text:p></table:table-cell></table:table-row>
 <table:table-row><table:table-cell><text:p>Кев</text:p></table:table-cell>
  <table:table-cell><text:p>Болота</text:p></table:table-cell><table:table-cell table:number-columns-repeated="16384"/></table:table-row>
 <table:table-row table:number-rows-repeated="1048000"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>
</table:table></office:spreadsheet></office:body></office:document-content>"""


def make_zip(path, files):
    import zipfile
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return path


def make_xlsx(path):
    main = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    rel = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    return make_zip(path, {
        "xl/workbook.xml": f'<workbook {main} {rel}><sheets><sheet name="Народы" sheetId="1" r:id="rId1"/>'
                           '</sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
                                      'relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml" '
                                      'Type="x"/></Relationships>',
        "xl/sharedStrings.xml": f'<sst {main}><si><t>Название</t></si><si><t>Сколько живут</t></si>'
                                '<si><t>Ния</t></si></sst>',
        "xl/worksheets/sheet1.xml": f'<worksheet {main}><sheetData>'
                                    '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c></row>'
                                    '<row r="2"><c r="A2" t="s"><v>2</v></c><c r="B2"><v>300.0</v></c></row>'
                                    '<row r="3"><c r="A3" t="inlineStr"><is><t>Давнь</t></is></c></row>'
                                    '</sheetData></worksheet>'})


class MoreFormatsTest(Sandbox):
    def test_odt(self):
        path = make_zip(self.tmp / "книга.odt", {"content.xml": ODT_CONTENT})
        paragraphs = I.read_document(path)
        self.assertIn(("Глава 1. Туман", 1), paragraphs)
        self.assertIn(("Обычный *курсив* и **жирный**.", 0), paragraphs)
        self.assertIn(("Строка", 0), paragraphs)
        self.assertEqual([name for name, _ in I.split_chapters(paragraphs, "книга")], ["Глава 1. Туман", "Глава 2"])

    def test_epub_and_fb2_from_our_own_export(self):
        from stungkir.export import export_epub, export_fb2
        chapters = [("Пролог", ["Первый *абзац*.", "Второй **абзац**."]), ("Глава 1", ["Текст главы."])]
        options = {"headings": True, "new_page": True, "layout": {}, "genre": "sf_fantasy"}
        export_epub(self.tmp / "к.epub", "Книга", "Автор", chapters, options)
        export_fb2(self.tmp / "к.fb2", "Книга", "Автор", chapters, options)
        for name in ("к.epub", "к.fb2"):
            paragraphs = I.read_document(self.tmp / name)
            split = I.split_chapters(paragraphs, "к")
            names = [n for n, _ in split]
            self.assertIn("Пролог", names, name)
            self.assertIn("Глава 1", names, name)
            body = dict(split)["Пролог"]
            self.assertIn("Первый *абзац*.", body, name)
            self.assertIn("Второй **абзац**.", body, name)

    def test_html(self):
        (self.tmp / "с.html").write_text("<html><head><title>x</title><style>p{}</style></head><body>"
                                         "<h1>Глава</h1><p>Раз <em>курсив</em>,<br>строка</p>"
                                         "<script>alert(1)</script></body></html>", encoding="utf-8", newline="\n")
        self.assertEqual(I.read_document(self.tmp / "с.html"),
                         [("Глава", 1), ("Раз *курсив*,", 0), ("строка", 0)])

    def test_tables(self):
        [(name, rows)] = I.xlsx_tables(make_xlsx(self.tmp / "т.xlsx"))
        self.assertEqual(name, "Народы")
        self.assertEqual(rows, [["Название", "Сколько живут"], ["Ния", "300"], ["Давнь", ""]])
        [(name, rows)] = I.ods_tables(make_zip(self.tmp / "т.ods", {"content.xml": ODS_CONTENT}))
        self.assertEqual((name, rows), ("Бестиарий", [["Существо", "Где водится"], ["Кев", "Болота"]]))


class ImportWindowTest(WindowCase):
    def test_add_files_of_every_kind(self):
        from PySide6.QtWidgets import QFileDialog
        from stungkir.textutil import parse_table
        w = self.window()
        category = C.ROOT / "Мир"
        category.mkdir()
        files = [make_zip(self.tmp / "книга.odt", {"content.xml": ODT_CONTENT}), make_xlsx(self.tmp / "т.xlsx")]
        (self.tmp / "скан.pdf").write_bytes(b"%PDF-1.4")
        files.append(self.tmp / "скан.pdf")
        self.patch("ask_choice", lambda *a, **k: "split")
        w.ask_category = lambda: category
        original = QFileDialog.getOpenFileNames
        QFileDialog.getOpenFileNames = staticmethod(lambda *a, **k: ([str(f) for f in files], ""))
        self.addCleanup(setattr, QFileDialog, "getOpenFileNames", original)
        w.import_files()
        names = sorted(p.name for p in S.files_in(category))
        self.assertEqual(names, ["Глава 1. Туман.txt", "Глава 2.txt", "т.csv"])
        header, rows = parse_table(S.read_text(category / "т.csv"))[:2]
        self.assertEqual(header, ["Название", "Сколько живут"])
        self.assertIn("PDF не открывается", self.messages[-1][1])


class WorldKitTest(WindowCase):
    def test_world_kit(self):
        from stungkir import worldkit
        from stungkir.textutil import parse_table
        w = self.window()
        category = C.ROOT / "Мир"
        category.mkdir()
        made = worldkit.create_kit(category, list(worldkit.KITS))
        self.assertEqual(len(made), len(worldkit.KITS))
        header = parse_table(S.read_text(category / "Бестиарий.csv"))[0]
        self.assertEqual(header[0], "Существо")
        self.assertTrue((category / "_Магия.txt").exists())
        self.assertTrue((category / "Хронология мира.timeline").exists())
        again = worldkit.create_kit(category, ["Бестиарий"])                 # занято - рядом, не поверх
        self.assertEqual(again[0].name, "Бестиарий (2).csv")
        S.save_text(category / "Бестиарий.csv", "Существо,Где водится\nКев,Болота\n")
        w.toggle_lore_category(category)
        w.rebuild_lore()
        names = sorted(t.name for t in w.lore.terms)
        self.assertIn("Кев", names)                                          # строки таблиц - термины
        self.assertNotIn("Магия", names)                                     # общие карточки - нет
        dialog = worldkit.WorldKitDialog(w, "Мир", True)
        self.assertEqual(len(dialog.chosen()), len(worldkit.KITS))
        self.assertFalse(dialog.lore.isVisible())
