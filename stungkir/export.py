# StungKir - писательский блокнот. Copyright (C) 2026 KiriFesTung
# SPDX-License-Identifier: GPL-3.0-or-later
"""Сборка книги: PDF, Word, EPUB и текст."""

import html
import re
import zipfile
from datetime import datetime
from pathlib import Path
from PySide6.QtCore import QMarginsF, QRectF, Qt
from PySide6.QtGui import (
    QAbstractTextDocumentLayout, QColor, QFont, QPainter, QPageLayout, QPageSize, QPalette, QTextDocument,
)

from . import config as C
from .textutil import MARKER_RE, clean_control, markup_runs, plain_markup, plural


class ChapterError(Exception):
    """Сбой при сборке конкретной главы."""


# размеры страниц в миллиметрах
PAGE_SIZES = {"A4": (210, 297), "A5": (148, 210), "145×205 (Ridero)": (145, 205), "6×9 дюймов": (152.4, 228.6)}
# оформление книги по умолчанию - как было до шаблонов
LAYOUT = {
    "page": "A4", "margins": (20, 15, 20, 25),      # сверху, справа, снизу, слева, мм
    "font_pdf": None, "font_docx": None, "size": 12.0, "spacing": 1.5,   # None - шрифты из config
    "indent": 12.5, "align": "justify",             # абзацный отступ, мм; justify или left
    "title_page": "book",                           # book - название посередине, manuscript - для издательства
    "header": False,                                # колонтитул «Автор · Название» с номером страницы
    "scene_break": "* * *", "contacts": "",
}
# готовые шаблоны: что меняется относительно LAYOUT, плюс формат и галочки
TEMPLATES = {
    "Обычная книга": {"format": "pdf"},
    "Рукопись для издательства": {
        "format": "docx", "page": "A4", "margins": (20, 20, 20, 25), "font_pdf": "Liberation Serif",
        "size": 12.0, "spacing": 1.5, "indent": 12.5, "align": "left", "title_page": "manuscript",
        "header": True, "scene_break": "* * *", "new_page": True, "headings": True},
    "Самиздат EPUB": {"format": "epub", "align": "justify", "indent": 7.0, "new_page": True, "headings": True},
    "Самиздат FB2": {"format": "fb2", "new_page": True, "headings": True},
    "Печать, книжный формат": {
        "format": "pdf", "page": "A5", "margins": (15, 15, 18, 20), "size": 10.5, "spacing": 1.25,
        "indent": 7.0, "align": "justify", "new_page": True, "headings": True},
}
SCENE_RE = re.compile(r"^\s*(?:(?:\*\s*){3,}|#{1,3}|(?:~\s*){3,}|(?:•\s*){3,})\s*$")
CHARS_PER_SHEET = 40000      # авторский лист
# заголовок внутри главы: «# Пролог», «## Часть вторая». Одна «#» без текста - разрыв сцены
HEADING_RE = re.compile(r"^\s*(#{1,3})\s+(\S.*?)\s*$")
# строка, целиком состоящая из пометки: «[ПРОЛОГ]»
MARKER_LINE_RE = re.compile(r"^\s*\[([^\[\]\n]{1,120})\]\s*$")


def layout_of(options):
    layout = dict(LAYOUT)
    layout.update(options.get("layout") or {})
    layout["font_pdf"] = layout["font_pdf"] or C.BOOK_FONT
    layout["font_docx"] = layout["font_docx"] or C.BOOK_FONT_DOCX
    return layout


def is_scene_break(text):
    return bool(SCENE_RE.match(text))


def heading_of(text):
    """(уровень, текст) для строки-заголовка «# …», иначе None."""
    m = HEADING_RE.match(text)
    return (len(m.group(1)), plain_markup(m.group(2))) if m else None


def marker_lines(text):
    """Строки, которые целиком - пометка в скобках: при сборке книги они бы пропали."""
    return [m.group(0).strip() for line in text.split("\n") for m in [MARKER_LINE_RE.match(line)] if m]


def book_volume(chapters):
    """(знаков с пробелами, слов) - для титула рукописи."""
    chars = words = 0
    for _name, paragraphs in chapters:
        for text in paragraphs:
            if is_scene_break(text):
                continue
            heading = heading_of(text)
            plain = heading[1] if heading else plain_markup(text)
            chars += len(plain)
            words += len(re.findall(r"\w+(?:[-'’]\w+)*", plain))
    return chars, words


def spaced(number):
    return f"{number:,}".replace(",", "\u00a0")


def volume_line(chapters):
    chars, words = book_volume(chapters)
    sheets = f"{chars / CHARS_PER_SHEET:.1f}".replace(".", ",")
    return (f"Объём: {spaced(chars)} {plural(chars, 'знак', 'знака', 'знаков')} с пробелами "
            f"({sheets} а. л.), {spaced(words)} {plural(words, 'слово', 'слова', 'слов')}")


def book_paragraphs(text, strip_markers, marker_headings=False):
    """Текст главы -> абзацы книги. marker_headings: строки вида «[ПРОЛОГ]» становятся
    заголовками «# ПРОЛОГ», а не вырезаются вместе с остальными пометками."""
    text = clean_control(text)
    if marker_headings:
        text = "\n".join(MARKER_LINE_RE.sub(lambda m: f"# {m.group(1).strip()}", line)
                          for line in text.split("\n"))
    if strip_markers:
        text = MARKER_RE.sub("", text)
    paragraphs = []
    for line in text.replace("\r\n", "\n").split("\n"):
        line = re.sub(r"[ \t]{2,}", " ", line).strip()
        if line:
            paragraphs.append(line)
    return paragraphs


def export_txt(target, title, author, chapters, options):
    parts = [title]
    if author:
        parts.append(author)
    for name, paragraphs in chapters:
        block = [name, ""] if options["headings"] else []
        layout = layout_of(options)
        block += [layout["scene_break"] if is_scene_break(p) else
                  heading_of(p)[1] if heading_of(p) else plain_markup(p) for p in paragraphs]
        parts.append("\n".join(block))
    target.write_text("\n\n\n".join(parts) + "\n", encoding="utf-8")


EPUB_CSS = """\
body { margin: 0 4%; line-height: 1.5; }
h1 { font-size: 1.5em; text-align: center; margin: 2.5em 0 1.5em; page-break-after: avoid; }
p { margin: 0; text-indent: INDENT; text-align: ALIGN; }
h1 + p, .first { text-indent: 0; }
p.scene { text-indent: 0; text-align: center; margin: 1em 0; }
h2, h3, h4 { font-size: 1.2em; text-align: center; margin: 1.6em 0 1em; page-break-after: avoid; }
.title { text-align: center; margin-top: 30%; }
.title h1 { font-size: 2em; margin: 0 0 0.6em; }
.title p { text-indent: 0; text-align: center; font-size: 1.2em; }
.cover { text-align: center; margin: 0; }
.cover img { max-width: 100%; max-height: 100%; }
"""


def epub_markup(text):
    out = []
    for piece, bold, italic in markup_runs(text):
        piece = html.escape(piece)
        if italic:
            piece = f"<em>{piece}</em>"
        if bold:
            piece = f"<strong>{piece}</strong>"
        out.append(piece)
    return "".join(out)


def export_epub(target, title, author, chapters, options):
    """EPUB 3 для читалок и самиздата: титул, главы отдельными файлами, оглавление.
    Оглавление продублировано в старом формате NCX, чтобы открывалось и на старых читалках."""
    import uuid
    from datetime import timezone
    title, author = clean_control(title), clean_control(author)
    layout = layout_of(options)
    esc = html.escape
    book_id = f"urn:uuid:{uuid.uuid4()}"
    modified = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

    def page(page_title, body):
        return ('<?xml version="1.0" encoding="utf-8"?>\n<!DOCTYPE html>\n'
                '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" '
                'xml:lang="ru" lang="ru">\n'
                f'<head>\n<meta charset="utf-8"/>\n<title>{esc(page_title)}</title>\n'
                '<link rel="stylesheet" type="text/css" href="style.css"/>\n</head>\n'
                f'<body>\n{body}\n</body>\n</html>\n')

    files = [("title.xhtml", title, page(title, '<section class="title" epub:type="titlepage">\n'
                                               f'<h1>{esc(title)}</h1>\n'
                                               + (f'<p>{esc(author)}</p>\n' if author else '')
                                               + '</section>'))]
    for number, (name, paragraphs) in enumerate(chapters, 1):
        name = clean_control(name)
        try:
            body = ['<section epub:type="chapter">']
            if options["headings"]:
                body.append(f"<h1>{esc(name)}</h1>")
            after_break = not options["headings"]
            for text in paragraphs:
                if is_scene_break(text):
                    body.append(f'<p class="scene">{esc(layout["scene_break"])}</p>')
                    after_break = True
                    continue
                heading = heading_of(text)
                if heading:
                    body.append(f"<h{heading[0] + 1}>{esc(heading[1])}</h{heading[0] + 1}>")
                    after_break = True
                    continue
                first = ' class="first"' if after_break else ""
                after_break = False
                body.append(f"<p{first}>{epub_markup(clean_control(text))}</p>")
            body.append("</section>")
        except Exception as error:
            raise ChapterError(f"глава «{name}»: {error}") from error
        files.append((f"chapter-{number:03}.xhtml", name, page(name, "\n".join(body))))

    nav_items = "\n".join(f'<li><a href="{f}">{esc(label)}</a></li>' for f, label, _ in files[1:])
    nav = page("Оглавление", '<nav epub:type="toc" id="toc">\n<h1>Оглавление</h1>\n'
                             f'<ol>\n{nav_items}\n</ol>\n</nav>')
    ncx_points = "\n".join(
        f'<navPoint id="p{i}" playOrder="{i}"><navLabel><text>{esc(label)}</text></navLabel>'
        f'<content src="{f}"/></navPoint>' for i, (f, label, _) in enumerate(files[1:], 1))
    ncx = ('<?xml version="1.0" encoding="utf-8"?>\n'
           '<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">\n'
           f'<head><meta name="dtb:uid" content="{book_id}"/></head>\n'
           f'<docTitle><text>{esc(title)}</text></docTitle>\n'
           f'<navMap>\n{ncx_points}\n</navMap>\n</ncx>\n')
    manifest = "\n".join(f'<item id="f{i}" href="{f}" media-type="application/xhtml+xml"/>'
                         for i, (f, _label, _) in enumerate(files))
    spine = "\n".join(f'<itemref idref="f{i}"/>' for i in range(len(files)))
    cover = options.get("cover")
    cover_meta = ""
    if cover:                                 # своя обложка категории: картинка и первая страница
        manifest = ('<item id="cover-image" href="cover.jpg" media-type="image/jpeg" properties="cover-image"/>\n'
                    '<item id="cover" href="cover.xhtml" media-type="application/xhtml+xml"/>\n' + manifest)
        spine = '<itemref idref="cover"/>\n' + spine
        cover_meta = '<meta name="cover" content="cover-image"/>\n'
    cover_page = page("Обложка", '<section epub:type="cover" class="cover">'
                                 f'<img src="cover.jpg" alt="{esc(title)}"/></section>')
    creator = f"<dc:creator>{esc(author)}</dc:creator>\n" if author else ""
    opf = ('<?xml version="1.0" encoding="utf-8"?>\n'
           '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid" xml:lang="ru">\n'
           '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
           f'<dc:identifier id="bookid">{book_id}</dc:identifier>\n'
           f'<dc:title>{esc(title)}</dc:title>\n<dc:language>ru</dc:language>\n{creator}'
           f'<meta property="dcterms:modified">{modified}</meta>\n{cover_meta}</metadata>\n'
           '<manifest>\n<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>\n'
           '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>\n'
           '<item id="css" href="style.css" media-type="text/css"/>\n'
           f'{manifest}\n</manifest>\n<spine toc="ncx">\n{spine}\n</spine>\n</package>\n')
    container = ('<?xml version="1.0" encoding="utf-8"?>\n'
                 '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
                 '<rootfiles><rootfile full-path="OEBPS/content.opf" '
                 'media-type="application/oebps-package+xml"/></rootfiles>\n</container>\n')
    with zipfile.ZipFile(target, "w") as book:
        # mimetype - первым и без сжатия, так требует формат
        book.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        book.writestr("META-INF/container.xml", container, zipfile.ZIP_DEFLATED)
        book.writestr("OEBPS/content.opf", opf, zipfile.ZIP_DEFLATED)
        book.writestr("OEBPS/nav.xhtml", nav, zipfile.ZIP_DEFLATED)
        book.writestr("OEBPS/toc.ncx", ncx, zipfile.ZIP_DEFLATED)
        css = (EPUB_CSS.replace("INDENT", f"{max(0.0, layout['indent']) / 5:.2f}em")
               .replace("ALIGN", "justify" if layout["align"] == "justify" else "left"))
        book.writestr("OEBPS/style.css", css, zipfile.ZIP_DEFLATED)
        for f, _label, content in files:
            book.writestr(f"OEBPS/{f}", content, zipfile.ZIP_DEFLATED)
        if cover:
            book.writestr("OEBPS/cover.jpg", cover, zipfile.ZIP_STORED)
            book.writestr("OEBPS/cover.xhtml", cover_page, zipfile.ZIP_DEFLATED)


FB2_GENRES = {"sf_fantasy": "Фэнтези", "sf": "Фантастика", "prose_contemporary": "Современная проза",
              "prose_classic": "Классическая проза", "love_contemporary": "Любовный роман",
              "detective": "Детектив", "thriller": "Триллер", "adventure": "Приключения",
              "child_prose": "Детская проза", "poetry": "Поэзия"}


def fb2_markup(text):
    out = []
    for piece, bold, italic in markup_runs(text):
        piece = html.escape(piece, quote=False)
        if italic:
            piece = f"<emphasis>{piece}</emphasis>"
        if bold:
            piece = f"<strong>{piece}</strong>"
        out.append(piece)
    return "".join(out)


def fb2_author(author):
    """«Имя Фамилия» -> имя и фамилия, одно слово (псевдоним) -> nickname."""
    esc = lambda s: html.escape(s, quote=False)      # noqa: E731
    words = author.split()
    if not words:
        return "<author><nickname>Автор</nickname></author>"
    if len(words) == 1:
        return f"<author><nickname>{esc(words[0])}</nickname></author>"
    middle = f"<middle-name>{esc(' '.join(words[1:-1]))}</middle-name>" if len(words) > 2 else ""
    return f"<author><first-name>{esc(words[0])}</first-name>{middle}<last-name>{esc(words[-1])}</last-name></author>"


def export_fb2(target, title, author, chapters, options):
    """FB2 - главный формат русских читалок и самиздат-площадок: один XML-файл в UTF-8."""
    import uuid
    from datetime import date
    from . import __version__
    title, author = clean_control(title), clean_control(author)
    layout = layout_of(options)
    esc = lambda s: html.escape(s, quote=False)      # noqa: E731
    genre = options.get("genre") or "prose_contemporary"
    today = date.today()
    body = [f"<title><p>{esc(title)}</p></title>"]
    for name, paragraphs in chapters:
        name = clean_control(name)
        try:
            section = ["<section>"]
            if options["headings"]:
                section.append(f"<title><p>{esc(name)}</p></title>")
            for text in paragraphs:
                text = clean_control(text)
                if is_scene_break(text):
                    section.append(f"<subtitle>{esc(layout['scene_break'])}</subtitle>")
                    continue
                heading = heading_of(text)
                if heading:
                    section.append(f"<subtitle>{esc(heading[1])}</subtitle>")
                    continue
                section.append(f"<p>{fb2_markup(text)}</p>" if text.strip() else "<empty-line/>")
            if len(section) == 1 or (len(section) == 2 and options["headings"]):
                section.append("<empty-line/>")          # пустая глава: секция всё равно должна что-то содержать
            section.append("</section>")
        except Exception as error:
            raise ChapterError(f"глава «{name}»: {error}") from error
        body.append("\n".join(section))
    nick = author.split()[0] if author.split() else "Автор"
    xml = ('<?xml version="1.0" encoding="UTF-8"?>\n'
           '<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0" '
           'xmlns:l="http://www.w3.org/1999/xlink">\n'
           "<description>\n<title-info>\n"
           f"<genre>{genre}</genre>\n{fb2_author(author)}\n<book-title>{esc(title)}</book-title>\n"
           "<lang>ru</lang>\n</title-info>\n<document-info>\n"
           f"<author><nickname>{esc(nick)}</nickname></author>\n"
           f"<program-used>{C.APP_NAME} {__version__}</program-used>\n"
           f'<date value="{today.isoformat()}">{today:%d.%m.%Y}</date>\n'
           f"<id>{uuid.uuid4()}</id>\n<version>1.0</version>\n</document-info>\n</description>\n"
           "<body>\n" + "\n".join(body) + "\n</body>\n</FictionBook>\n")
    Path(target).write_text(xml, encoding="utf-8")


def _pdf_font(family):
    """Шрифт для PDF: если нужного нет в системе, берём DejaVu Serif, он стоит всегда."""
    from PySide6.QtGui import QFontDatabase
    families = set(QFontDatabase.families())
    for name in (family, C.BOOK_FONT, "DejaVu Serif"):
        if name in families:
            return name
    return family


def export_pdf(target, title, author, chapters, options):
    """PDF по шаблону: размер страницы, поля, шрифт, интервал, абзацный отступ, титул, номера страниц."""
    from PySide6.QtCore import QSizeF
    from PySide6.QtPrintSupport import QPrinter

    layout = layout_of(options)
    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(str(target))
    width, height = PAGE_SIZES.get(layout["page"], PAGE_SIZES["A4"])
    if layout["page"] == "A4":
        printer.setPageSize(QPageSize(QPageSize.A4))
    elif layout["page"] == "A5":
        printer.setPageSize(QPageSize(QPageSize.A5))
    else:
        printer.setPageSize(QPageSize(QSizeF(width, height), QPageSize.Millimeter, layout["page"]))
    top, right, bottom, left = layout["margins"]
    printer.setPageMargins(QMarginsF(left, top, right, bottom), QPageLayout.Millimeter)
    printer.setDocName(title)
    printer.setCreator(C.APP_NAME)
    dpi = printer.resolution()
    font = _pdf_font(layout["font_pdf"])
    size = layout["size"]

    def mm(value):   # отступ первой строки Qt считает в точках принтера
        return round(value * dpi / 25.4)

    def gap(value):  # а поля между абзацами - в экранных пикселях
        return round(value * 96 / 25.4)

    def paragraph_html(text, first):
        if is_scene_break(text):
            return f'<p class="scene">{html.escape(layout["scene_break"])}</p>'
        heading = heading_of(text)
        if heading:
            return f'<h3>{html.escape(heading[1])}</h3>'
        out = []
        for piece, bold, italic in markup_runs(text):
            piece = html.escape(piece)
            if bold:
                piece = f"<b>{piece}</b>"
            if italic:
                piece = f"<i>{piece}</i>"
            out.append(piece)
        return ('<p class="first">' if first else "<p>") + "".join(out) + "</p>"

    if layout["title_page"] == "manuscript":
        contacts = "<br>".join(html.escape(x) for x in [author] + layout["contacts"].split("\n") if x.strip())
        body = [f'<table width="100%"><tr><td class="left">{contacts}</td>'
                f'<td class="right">{html.escape(volume_line(chapters))}</td></tr></table>',
                '<p class="book">&nbsp;</p>' * 8, f'<p class="book">{html.escape(title)}</p>']
    else:
        # у первого абзаца Qt игнорирует отступ сверху, поэтому опускаем название пустыми строками
        body = ['<p class="book">&nbsp;</p>' * 5, f'<p class="book">{html.escape(title)}</p>']
    if author:
        body.append(f'<p class="author">{html.escape(author)}</p>')
    for i, (name, paragraphs) in enumerate(chapters):
        if options["new_page"] or i == 0:
            body.append('<div style="page-break-before: always"></div>')
        if options["headings"]:
            body.append(f"<h2>{html.escape(name)}</h2>")
        first = options["headings"]
        for text in paragraphs:
            body.append(paragraph_html(text, first))
            first = is_scene_break(text) or bool(heading_of(text))

    document = QTextDocument()
    document.documentLayout().setPaintDevice(printer)
    base = QFont(font)
    base.setPointSizeF(size)
    document.setDefaultFont(base)
    document.setDocumentMargin(0)
    align = "justify" if layout["align"] == "justify" else "left"
    document.setDefaultStyleSheet(
        f"p {{ text-indent: {mm(layout['indent'])}px; margin: 0; line-height: {round(layout['spacing'] * 100)}%;"
        f" text-align: {align}; }}"
        "p.first { text-indent: 0px; }"
        f"p.scene {{ text-indent: 0px; text-align: center; margin-top: {gap(4)}px; margin-bottom: {gap(4)}px; }}"
        f"h2 {{ text-align: center; font-size: {size + 4:g}pt; margin-top: {gap(10)}px; margin-bottom: {gap(10)}px; }}"
        f"h3 {{ text-align: center; font-size: {size + 2:g}pt; margin-top: {gap(7)}px; margin-bottom: {gap(5)}px; }}"
        f".book {{ text-align: center; text-indent: 0px; font-size: {size * 2.1:.0f}pt; font-weight: bold; }}"
        f".author {{ text-align: center; text-indent: 0px; font-size: {size + 2:g}pt; margin-top: {gap(6)}px; }}"
        f"td.left {{ text-align: left; font-size: {size - 1:g}pt; }}"
        f"td.right {{ text-align: right; font-size: {size - 1:g}pt; }}")
    document.setHtml("".join(body))

    area = printer.pageRect(QPrinter.DevicePixel)
    footer = mm(10)
    header = mm(9) if layout["header"] else 0
    page_height = area.height() - footer - header
    document.setPageSize(QSizeF(area.width(), page_height))
    painter = QPainter(printer)
    if options.get("cover"):                  # своя обложка категории - первой страницей, во весь лист
        from PySide6.QtGui import QImage
        image = QImage.fromData(options["cover"])
        if not image.isNull():
            full = printer.paperRect(QPrinter.DevicePixel)
            page_area = printer.pageRect(QPrinter.DevicePixel)
            fitted = image.size().scaled(full.size().toSize(), Qt.KeepAspectRatio)
            spot = QRectF(full.x() - page_area.x() + (full.width() - fitted.width()) / 2,
                          full.y() - page_area.y() + (full.height() - fitted.height()) / 2,
                          fitted.width(), fitted.height())
            painter.drawImage(spot, image)
            printer.newPage()
    small = QFont(font)
    small.setPointSizeF(max(7.0, size - 2))
    running = " · ".join(x for x in (author, title) if x)
    for page in range(document.pageCount()):
        if page:
            printer.newPage()
        painter.save()
        painter.translate(0, header - page * page_height)
        # обрезаем на пару точек принтера выше границы: иначе при печати может проступить
        # волосок верхушки строки, которая уже ушла на следующую страницу
        # рисуем с чёрной палитрой: иначе Qt берёт цвет текста из темы программы (зелёный на белом)
        context = QAbstractTextDocumentLayout.PaintContext()
        context.palette.setColor(QPalette.Text, QColor(0, 0, 0))
        context.palette.setColor(QPalette.WindowText, QColor(0, 0, 0))
        context.clip = QRectF(0, page * page_height, area.width(), page_height - 2)
        painter.setClipRect(context.clip)
        document.documentLayout().draw(painter, context)
        painter.restore()
        if not page:
            continue   # на титульной странице номер и колонтитул не ставим
        painter.setFont(small)
        painter.setPen(QColor(0, 0, 0))
        if header:
            painter.drawText(QRectF(0, 0, area.width(), header), Qt.AlignLeft | Qt.AlignTop, running)
            painter.drawText(QRectF(0, 0, area.width(), header), Qt.AlignRight | Qt.AlignTop, str(page + 1))
        else:
            painter.drawText(QRectF(0, header + page_height, area.width(), footer),
                             Qt.AlignHCenter | Qt.AlignBottom, str(page + 1))
    painter.end()


def export_docx(target, title, author, chapters, options):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Mm, Pt, RGBColor

    layout = layout_of(options)
    font, size = layout["font_docx"], layout["size"]
    book = Document()
    normal = book.styles["Normal"]
    normal.font.name = font
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    normal.font.size = Pt(size)
    normal.paragraph_format.first_line_indent = Mm(layout["indent"])
    normal.paragraph_format.line_spacing = layout["spacing"]
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = (WD_ALIGN_PARAGRAPH.JUSTIFY if layout["align"] == "justify"
                                         else WD_ALIGN_PARAGRAPH.LEFT)
    heading = book.styles["Heading 1"]
    heading.font.name = font
    heading.font.size = Pt(size + 4)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Cm(0)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(18)
    width, height = PAGE_SIZES.get(layout["page"], PAGE_SIZES["A4"])
    top, right, bottom, left = layout["margins"]

    def page_field(run):
        for kind, value in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                element = OxmlElement("w:fldChar")
                element.set(qn("w:fldCharType"), kind)
            else:
                element = OxmlElement("w:instrText")
                element.set(qn("xml:space"), "preserve")
                element.text = value
            run._r.append(element)

    for section in book.sections:
        section.page_width, section.page_height = Mm(width), Mm(height)
        section.top_margin, section.right_margin = Mm(top), Mm(right)
        section.bottom_margin, section.left_margin = Mm(bottom), Mm(left)
        section.different_first_page_header_footer = True   # на титуле без номера
        if layout["header"]:
            line = section.header.paragraphs[0]
            line.paragraph_format.first_line_indent = Cm(0)
            line.paragraph_format.tab_stops.add_tab_stop(Mm(width - left - right), WD_TAB_ALIGNMENT.RIGHT)
            run = line.add_run(clean_control(" · ".join(x for x in (author, title) if x)) + "\t")
            run.font.size = Pt(size - 2)
            page_field(line.add_run())
        else:
            footer = section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer.paragraph_format.first_line_indent = Cm(0)
            page_field(footer.add_run())

    def centered(text, points, bold=False, before=0):
        line = book.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.paragraph_format.first_line_indent = Cm(0)
        line.paragraph_format.space_before = Pt(before)
        run = line.add_run(clean_control(text))
        run.bold = bold
        run.font.size = Pt(points)
        return line

    if layout["title_page"] == "manuscript":
        block = book.add_paragraph()
        block.paragraph_format.first_line_indent = Cm(0)
        block.paragraph_format.line_spacing = 1.0
        block.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        block.add_run("\n".join(clean_control(x) for x in [author] + layout["contacts"].split("\n") if x.strip()))
        volume = book.add_paragraph()
        volume.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        volume.paragraph_format.first_line_indent = Cm(0)
        volume.add_run(volume_line(chapters)).font.size = Pt(size - 1)
        centered(title, size * 2, bold=True, before=180)
        if author:
            centered(author, size + 2, before=12)
    else:
        centered(title, 26, bold=True, before=220)
        if author:
            centered(author, 14, before=12)

    for i, (name, paragraphs) in enumerate(chapters):
        try:
            if options["new_page"] or i == 0:
                book.add_page_break()
            if options["headings"]:
                book.add_heading(clean_control(name), level=1)
            first = options["headings"]
            for text in paragraphs:
                if is_scene_break(text):
                    line = centered(layout["scene_break"], size)
                    line.paragraph_format.space_before = Pt(size)
                    line.paragraph_format.space_after = Pt(size)
                    first = True
                    continue
                heading = heading_of(text)
                if heading:
                    sub = book.add_heading(clean_control(heading[1]), level=min(3, heading[0] + 1))
                    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in sub.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        run.font.name = font
                    first = True
                    continue
                paragraph = book.add_paragraph()
                if first:
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                first = False
                for piece, bold, italic in markup_runs(clean_control(text)):
                    run = paragraph.add_run(piece)
                    run.bold = bold
                    run.italic = italic
        except Exception as error:
            raise ChapterError(f"глава «{name}»: {error}") from error
    book.save(str(target))


def docx_available():
    import importlib.util
    return importlib.util.find_spec("docx") is not None
